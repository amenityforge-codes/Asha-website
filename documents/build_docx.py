"""Generate a publication-ready Word document for Clause 3."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x2D, 0x3A, 0x2F)
MARK = RGBColor(0xD9, 0x6B, 0x27)
INK = RGBColor(0x1A, 0x19, 0x17)
MUTED = RGBColor(0x6E, 0x6A, 0x63)

OUT = Path(__file__).with_name("Aims-Mission-and-Objectives-of-the-Association.docx")


def set_run(run, *, size=11, bold=False, italic=False, color=INK, all_caps=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if all_caps:
        run.font.all_caps = True


def add_centered(doc, text, *, size=11, bold=False, italic=False, color=INK, space_after=6, all_caps=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color, all_caps=all_caps)
    return p


def add_heading_clause(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=12, bold=True, color=ACCENT)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11.5, bold=True, color=INK)
    return p


def add_body(doc, text, *, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.75)
    run = p.add_run(text)
    set_run(run, size=11)
    return p


def add_clause(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(2.7)
    p.paragraph_format.first_line_indent = Cm(-2.7)
    num = p.add_run(f"{number}\t")
    set_run(num, size=11, bold=True, color=MARK)
    body = p.add_run(text)
    set_run(body, size=11)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(2.7), WD_TAB_ALIGNMENT.LEFT)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = INK

    add_centered(
        doc,
        "MEMORANDUM OF ASSOCIATION / RULES & REGULATIONS",
        size=10,
        bold=True,
        color=ACCENT,
        space_after=4,
        all_caps=True,
    )
    add_centered(
        doc,
        "AIMS, MISSION AND OBJECTIVES OF THE ASSOCIATION",
        size=16,
        bold=True,
        space_after=4,
        all_caps=True,
    )
    add_centered(
        doc,
        "A non-profit society representing three-star and above hotels across Andhra Pradesh",
        size=11,
        italic=True,
        color=MUTED,
        space_after=18,
    )

    add_heading_clause(doc, "3. AIMS, MISSION AND OBJECTIVES OF THE ASSOCIATION")

    add_subheading(doc, "3.1 Establishment and Core Purpose")
    add_body(
        doc,
        "The Association is established as a non-profit society to promote the growth, collaboration, representation, professional excellence, and sustainable development of hotels classified as three-star and above across all twenty-six (26) districts of the State of Andhra Pradesh.",
    )
    add_body(
        doc,
        "The Association shall serve as a unified institutional platform representing the hospitality industry and facilitating cooperation among member hotels, government departments, Andhra Pradesh Tourism authorities, local authorities, industry stakeholders, educational and training institutions, national and international hospitality organisations, and other relevant public and private institutions.",
    )
    add_body(
        doc,
        "All activities of the Association shall be carried out on a non-profit basis, solely in furtherance of the Mission, Aims, and Objectives set out herein, and not for the private profit or personal gain of any member.",
    )

    add_subheading(doc, "3.2 Mission")
    add_body(
        doc,
        "To represent, strengthen, and advance three-star and above hotels across Andhra Pradesh by promoting hospitality excellence, high standards of service, hygiene, food safety, sustainability, and guest and operational safety; advocating tourism-friendly public policy; fostering professional development; and positioning Andhra Pradesh as a premier tourism and hospitality destination.",
    )

    add_subheading(doc, "3.3 Aims")
    add_body(
        doc,
        "The Aims of the Association, which give effect to the Mission, are as follows:",
        first_line=False,
    )
    add_clause(
        doc,
        "3.3.1",
        "To support and represent member hotels in matters of public policy, investment, tourism development, industry regulation, and public engagement.",
    )
    add_clause(
        doc,
        "3.3.2",
        "To provide a unified platform for structured collaboration among government authorities, private stakeholders, educational institutions, and the hospitality industry.",
    )
    add_clause(
        doc,
        "3.3.3",
        "To uphold and promote benchmarks of service quality, hygiene, food safety, guest and operational safety, professionalism, and sustainability.",
    )
    add_clause(
        doc,
        "3.3.4",
        "To promote Andhra Pradesh as a premier tourism and hospitality destination through advocacy, training, capacity-building, recognition, and public engagement.",
    )
    add_clause(
        doc,
        "3.3.5",
        "To foster the growth, collaboration, professional excellence, and sustainable development of the organised hotel sector throughout the State.",
    )

    add_subheading(doc, "3.4 Non-Profit Character")
    add_clause(doc, "3.4.1", "The Association is constituted and shall function exclusively on a non-profit basis.")
    add_clause(
        doc,
        "3.4.2",
        "The income, property, funds, and assets of the Association shall be applied solely towards the attainment of its Aims and Objectives.",
    )
    add_clause(
        doc,
        "3.4.3",
        "No portion of the income or property of the Association shall be paid, transferred, or distributed, directly or indirectly, by way of dividend, bonus, profit, or otherwise, to any member, office-bearer, or other person, except as reasonable remuneration, reimbursement, or consideration for services actually rendered, or expenses actually incurred, in furtherance of the Association’s Objectives, and in accordance with its governing documents and applicable laws and regulations.",
    )
    add_clause(
        doc,
        "3.4.4",
        "Any activity of a commercial, trading, or revenue-generating nature that the Association may lawfully undertake shall be incidental or ancillary to the Objectives herein, and the proceeds thereof shall be applied exclusively to the Association’s non-profit purposes.",
    )

    add_subheading(doc, "3.5 Objectives")
    add_body(
        doc,
        "In furtherance of the Mission and Aims, and not otherwise, the Association shall have the following Objectives:",
        first_line=False,
    )

    sections = [
        (
            "3.5.1 Representation and Advocacy",
            [
                ("3.5.1.1", "To represent the collective interests and voice of three-star and above hotels across all twenty-six (26) districts of Andhra Pradesh."),
                ("3.5.1.2", "To represent member hotels before government departments, tourism authorities, statutory bodies, regulators, local authorities, and other policy-making institutions."),
                ("3.5.1.3", "To liaise and collaborate with government departments, public authorities, and private stakeholders on matters affecting the hospitality and tourism sector."),
                ("3.5.1.4", "To advocate policies, measures, and programmes that support tourism, hospitality, investment, employment, infrastructure, and the orderly growth of the industry."),
                ("3.5.1.5", "To provide a structured and continuing platform for communication, consultation, and cooperation between the government and the hospitality industry."),
                ("3.5.1.6", "To ensure, for the purpose of effective government liaison and institutional coordination, that either the President or the General Secretary of the Association is ordinarily based in Vijayawada or Amaravati."),
            ],
        ),
        (
            "3.5.2 Quality, Standards and Professional Excellence",
            [
                ("3.5.2.1", "To promote high standards of hospitality service among member hotels and across the organised hotel sector in the State."),
                ("3.5.2.2", "To encourage excellence in hygiene, sanitation, food safety, guest safety, security, sustainability, and professional conduct."),
                ("3.5.2.3", "To establish, adopt, or promote appropriate industry benchmarks, codes of practice, and best practices, consistent with applicable laws and regulations."),
                ("3.5.2.4", "To facilitate knowledge-sharing, peer learning, and benchmarking among member hotels."),
                ("3.5.2.5", "To encourage responsible, ethical, and environmentally sustainable hospitality practices."),
            ],
        ),
        (
            "3.5.3 Member Development, Training and Recognition",
            [
                ("3.5.3.1", "To organise training programmes, seminars, conferences, workshops, and other professional development programmes for members, their employees, and other hospitality professionals, as appropriate."),
                ("3.5.3.2", "To provide or facilitate training relating to hotel operations, hospitality management, marketing, technology, statutory and regulatory compliance, sustainability, food safety, innovation, and human resource development."),
                ("3.5.3.3", "To encourage continuous professional development of persons employed in the hospitality sector."),
                ("3.5.3.4", "To recognise outstanding hotels, employees, professionals, and industry initiatives through awards and other appropriate forms of recognition."),
                ("3.5.3.5", "To provide opportunities for members to showcase their achievements, innovations, and best practices."),
            ],
        ),
        (
            "3.5.4 Crisis Management and Member Support",
            [
                ("3.5.4.1", "To represent and support members during natural disasters, pandemics, public emergencies, regulatory changes, economic disruptions, and other crises affecting the hospitality sector."),
                ("3.5.4.2", "To coordinate collective industry responses, where appropriate and lawful, in the interest of members and the public."),
                ("3.5.4.3", "To facilitate communication among member hotels, government departments, relief agencies, and other relevant institutions during emergencies."),
                ("3.5.4.4", "To support members in obtaining and sharing relevant information, guidance, and resources during such periods."),
            ],
        ),
        (
            "3.5.5 Tourism, Industry Promotion and Public Engagement",
            [
                ("3.5.5.1", "To promote Andhra Pradesh as a premier tourism and hospitality destination."),
                ("3.5.5.2", "To promote the State’s hospitality sector through national and international platforms."),
                ("3.5.5.3", "To participate in tourism fairs, exhibitions, conferences, trade events, and industry forums."),
                ("3.5.5.4", "To facilitate media engagement and public-awareness campaigns concerning hospitality, tourism, and related matters of common interest."),
                ("3.5.5.5", "To promote the positive image of member hotels and public understanding of their contribution to tourism, employment, investment, and the local and State economy."),
                ("3.5.5.6", "To publish or disseminate appropriate promotional and educational materials relating to hospitality and tourism."),
            ],
        ),
        (
            "3.5.6 Strategic Institutional Affiliations",
            [
                ("3.5.6.1", "To obtain and maintain memberships, affiliations, partnerships, and collaborations with relevant State, national, and international hospitality, tourism, trade, and professional organisations, subject to applicable laws and regulations and consistent with the Association’s non-profit character."),
                ("3.5.6.2", "To participate in federations, alliances, councils, chambers, and other appropriate institutions."),
                ("3.5.6.3", "To use such affiliations to advance the legitimate interests of members and to strengthen the Association’s institutional network."),
                ("3.5.6.4", "To ensure that all affiliations, partnerships, and collaborations remain consistent with the Association’s non-profit character and Objectives."),
            ],
        ),
        (
            "3.5.7 Educational, Institutional and Industry Development",
            [
                ("3.5.7.1", "To establish, support, or maintain, wherever appropriate and subject to applicable laws and regulations, charitable institutions, reading rooms, libraries, recreation facilities, and other useful activities for the benefit of members or the public."),
                ("3.5.7.2", "To establish or support hospitality, catering, hotel-management, and vocational training institutions, subject to applicable laws and regulations."),
                ("3.5.7.3", "To promote education, skill development, employment, and career opportunities within the hospitality sector."),
                ("3.5.7.4", "To establish or support food-quality, hygiene, chemical-analysis, or other testing laboratories for the benefit of member hotels, subject to applicable laws and regulations and to such licences, accreditations, and permissions as may be required."),
                ("3.5.7.5", "To promote research, knowledge-sharing, innovation, and technological development in hospitality and allied fields."),
            ],
        ),
        (
            "3.5.8 Publications, Communication and Publicity",
            [
                ("3.5.8.1", "To publish journals, newsletters, bulletins, periodicals, magazines, reports, research materials, and other publications in print, electronic, or other lawful form."),
                ("3.5.8.2", "To disseminate information concerning hospitality, tourism, regulations, standards, policies, and industry developments."),
                ("3.5.8.3", "To organise press conferences, media interactions, publicity programmes, and public-awareness initiatives on matters of common interest."),
                ("3.5.8.4", "To maintain effective communication channels with members, government authorities, industry stakeholders, and the public."),
            ],
        ),
        (
            "3.5.9 Exhibitions, Conferences and Industry Events",
            [
                ("3.5.9.1", "To organise or participate in exhibitions relating to food products, hotel equipment, technology, raw materials, services, and other hospitality-related products and services."),
                ("3.5.9.2", "To organise conferences, meetings, seminars, conventions, workshops, and industry forums."),
                ("3.5.9.3", "To provide platforms for discussion of policy, regulation, technology, investment, operations, tourism, and other matters affecting the hospitality industry."),
            ],
        ),
        (
            "3.5.10 Regional and Branch Development",
            [
                ("3.5.10.1", "To establish branch offices, regional offices, chapters, or other representative offices in different parts of Andhra Pradesh, as may be required for the effective functioning of the Association, subject to the Association’s governing documents and applicable laws and regulations."),
                ("3.5.10.2", "To ensure effective representation and participation of hotels from different regions and districts of the State."),
                ("3.5.10.3", "To promote balanced development and collaboration across the State’s hospitality sector."),
            ],
        ),
        (
            "3.5.11 Property, Assets and Financial Resources",
            [
                ("3.5.11.1", "To acquire, lease, develop, hold, manage, maintain, or otherwise deal with movable and immovable property required for achieving the Association’s Objectives, subject to applicable laws and regulations and to the Association’s governing documents."),
                ("3.5.11.2", "To establish offices, training facilities, meeting spaces, educational facilities, laboratories, libraries, or other infrastructure necessary or conducive to the Association’s Objectives, subject to applicable laws and regulations."),
                ("3.5.11.3", "To invest or deploy the funds of the Association only in accordance with applicable laws and regulations and solely in furtherance of the Association’s non-profit Objectives."),
                ("3.5.11.4", "To raise funds through lawful means, including membership fees, grants, donations, contributions, sponsorships, subscriptions, and other permissible sources, provided that such funds are applied exclusively towards the Association’s Objectives."),
                ("3.5.11.5", "To borrow funds, where legally permissible and necessary for achieving the Objectives of the Association, subject to applicable laws and regulations and to the Association’s governing documents."),
                ("3.5.11.6", "To manage, lease, mortgage, transfer, or otherwise deal with the property of the Association in accordance with applicable laws and regulations and the Association’s governing documents, and solely in furtherance of its Objectives."),
            ],
        ),
        (
            "3.5.12 General and Incidental Objectives",
            [
                ("3.5.12.1", "To undertake any lawful activity that is incidental, ancillary, or conducive to the attainment of the Association’s Aims and Objectives."),
                ("3.5.12.2", "To enter into lawful agreements, partnerships, collaborations, or other arrangements with public or private institutions in furtherance of the Association’s Objectives."),
                ("3.5.12.3", "To undertake activities relating to trade, commerce, public relations, education, tourism, hospitality, research, training, or industry development, where such activities directly support the Association’s stated Objectives, remain consistent with its non-profit character, and are subject to applicable laws and regulations."),
                ("3.5.12.4", "To take all other lawful measures necessary, expedient, or beneficial for achieving the Aims and Objectives of the Association, provided that no such measure shall confer private profit upon any member or be inconsistent with the Association’s non-profit character."),
            ],
        ),
    ]

    for title, clauses in sections:
        add_subheading(doc, title)
        for number, text in clauses:
            add_clause(doc, number, text)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.line_spacing = 1.15
    label = note.add_run("For insertion: ")
    set_run(label, size=10, bold=True, italic=True, color=MUTED)
    rest = note.add_run(
        "Insert the Association’s legal name before filing. Clause 3.5.1.6 should also be restated in the Rules governing office-bearers. A matching dissolution clause should provide that surplus assets, if any, shall vest in a similarly constituted non-profit body and not in the members. This document is institutional drafting and is not legal advice; have counsel licensed in Andhra Pradesh review it against the statute of registration and any intended tax-exemption application."
    )
    set_run(rest, size=10, italic=True, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Clause 3 - Aims, Mission and Objectives  |  Non-profit society  |  Andhra Pradesh")
    set_run(run, size=9, italic=True, color=MUTED)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
